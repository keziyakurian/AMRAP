from pptx import Presentation
from backend.config import config
import os

def write_ppt():
    print("Generating PPT Deliverable (Step 9)...")
    
    prs = Presentation()
    
    # Title Slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "AMRAP Analysis Report"
    subtitle.text = "Automated Generated Insight"
    
    # Example: Add Chart if exists
    # We would loop through charts/ dir
    
    output_file = os.path.join(config.OUTPUT_PATH, "final.pptx")
    prs.save(output_file)
    print(f"PPT saved to {output_file}")
